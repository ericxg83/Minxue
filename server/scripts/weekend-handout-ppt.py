#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
周末班错题课件 · 幻灯片生成器（PPTX + PDF）
============================================================================
消费 weekend-handout.mjs 产出的 <base>.slides.json，生成：

  1. <title>.pptx   —— 16:9，**每页一道题**；顶部信息条含「日期 / 难度 / 几人错 / 哪些学生」；
                       参考答案为一个圆角框，**默认不显示，放映时单击才淡入**（一键弹出答案）。
                       讲题要点、学生错答、原卷图链接写进「演讲者备注」。
  2. <title>.pdf    —— 同内容、同版式的投屏/打印备份（PDF 无法做「点击浮现」，答案直接印出）。

只读：不连数据库，只读 slides.json 与 OSS 图片（公开可读）。

用法：
  python server/scripts/weekend-handout-ppt.py deliverables/xxx.slides.json
  python server/scripts/weekend-handout-ppt.py xxx.slides.json --out <dir> --no-pdf
  python server/scripts/weekend-handout-ppt.py xxx.slides.json --no-animation    # 不出动画（答案直接可见）
  python server/scripts/weekend-handout-ppt.py xxx.slides.json --page-image     # 无配图时退贴整页原卷（默认关）
  # [2026-09-21] --wb-image 已下线（整题裁片 = 配图 B，写入侧已停；slides.json 不再下发 wbImage）

配图口径（2026-09-17 定）：
  **只贴 `questions.geometry_image_url` 系统裁片**；投屏绝不放整页原卷扫描
  （乱、带学生手写、缩小后看不清）。原卷有图但系统没裁出配图的题，
  只在页面上给一条提示，原卷链接留在「演讲者备注」里，由老师自己准备图形。
"""
import argparse
import concurrent.futures
import io
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.request
from pathlib import Path

# ────────────────────────────── 常量 ──────────────────────────────
SLIDE_W = 13.333
SLIDE_H = 7.5
M = 0.46                      # 左右页边距
HDR_Y = 0.30
HDR_H = 0.60
RULE_Y = 1.00
BODY_Y = 1.16
ANS_H = 1.44
ANS_Y = SLIDE_H - 0.36 - ANS_H
BODY_H = ANS_Y - 0.14 - BODY_Y
FIG_W = 4.75                  # 题图栏宽度
GUTTER = 0.28

FONT = "微软雅黑"

INK = (0x1B, 0x1F, 0x24)
INK2 = (0x4B, 0x55, 0x63)
INK3 = (0x8B, 0x94, 0x9E)
LINE = (0xE3, 0xE6, 0xEA)
BG_SOFT = (0xF7, 0xF8, 0xFA)
WHITE = (0xFF, 0xFF, 0xFF)

TIER_COLOR = {
    "basic": (0x0F, 0x7B, 0x4F),
    "medium": (0x1E, 0x6F, 0xD9),
    "hard": (0xC2, 0x41, 0x0C),
    "unknown": (0x6B, 0x72, 0x80),
}
TIER_CHIP_BG = {
    "basic": (0xE9, 0xF7, 0xF0),
    "medium": (0xEA, 0xF2, 0xFE),
    "hard": (0xFD, 0xF0, 0xE9),
    "unknown": (0xF1, 0xF3, 0xF5),
}
ANS_BG = (0xE9, 0xF7, 0xF0)
ANS_LINE = (0xC7, 0xEA, 0xD9)
ANS_INK = (0x0F, 0x7B, 0x4F)
WARN_BG = (0xFF, 0xFB, 0xEB)
WARN_LINE = (0xFD, 0xE6, 0x8A)
WARN_INK = (0xB4, 0x53, 0x09)

# ── 新版视觉（2026-09-17 重设计）──
PAGE_BG = (0xF2, 0xF4, 0xFB)      # 页面底：冷调浅灰蓝
CARD_LINE = (0xE5, 0xE9, 0xF4)    # 卡片描边
ACCENT = (0x4C, 0x51, 0xE0)       # 主色：靛蓝
GRAD1 = (0x4A, 0x4E, 0xE4)        # 封面/分节渐变起
GRAD2 = (0x7C, 0x38, 0xC9)        # 封面/分节渐变止
COND_BG = (0xEE, 0xF1, 0xFE)      # 公共题干底色
COND_BAR = (0xB9, 0xC2, 0xF5)     # 公共题干左条
COND_INK = (0x3F, 0x47, 0x8F)     # 公共题干文字
DOT_OFF = (0xDD, 0xE2, 0xF0)      # 难度点（空）

FIG_HINT = re.compile(r"如图|如下图|右图|左图|所示|图象|图像|图形|图中|看图|下列图形|统计图|示意图|网格")

CHROME_CANDIDATES = [
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
]


# ────────────────────────────── 数学文本规范化 ──────────────────────────────
def norm_math(s: str) -> str:
    """纯文本版（用于测量、备注）。上下标以 ^x / _x 形式保留。"""
    return "".join(t for t, _ in math_segments(s))


_MATH_MARK = re.compile(r"(\^|_)\{([^{}]*)\}|(\^|_)([0-9A-Za-z+\-=()√])")


def math_segments(s: str):
    """
    把题干/答案切成 [(text, kind)]，kind ∈ 'n' | 'sup' | 'sub'。

    ⚠ 不要用 Unicode 上下标字符（²⁶⁹ ₐₙ…）：微软雅黑缺大部分字形，
    在 PowerPoint 里会变成豆腐块。渲染端各自用「真上下标」实现：
    PPT 用 a:rPr/@baseline，HTML 用 <sup>/<sub>。
    """
    if not s:
        return []
    t = str(s).replace("\r\n", "\n").replace("\r", "\n").replace("$$", "").replace("$", "")
    t = re.sub(r"\\frac\s*\{([^{}]*)\}\s*\{([^{}]*)\}",
               lambda m: f"{m.group(1)}/{m.group(2)}", t)
    for _ in range(2):
        t = re.sub(r"\\frac\s*\{([^{}]*)\}\s*\{([^{}]*)\}",
                   lambda m: f"({m.group(1)})/({m.group(2)})", t)
    t = re.sub(r"\\sqrt\s*\{([^{}]*)\}", r"√(\1)", t)
    t = re.sub(r"\\sqrt\s*([A-Za-z0-9])", r"√\1", t)
    reps = {
        r"\times": "×", r"\div": "÷", r"\leq": "≤", r"\le": "≤", r"\geq": "≥", r"\ge": "≥",
        r"\neq": "≠", r"\ne": "≠", r"\pm": "±", r"\cdot": "·", r"\ldots": "…", r"\dots": "…",
        r"\angle": "∠", r"\triangle": "△", r"\circ": "°", r"\parallel": "∥", r"\perp": "⊥",
        r"\rightarrow": "→", r"\Rightarrow": "⇒", r"\because": "∵", r"\therefore": "∴",
        r"\left": "", r"\right": "", r"\,": " ", r"\;": " ", r"\!": "", r"\ ": " ",
        r"\%": "%", r"\#": "#", r"\&": "&",
    }
    for k, v in reps.items():
        t = t.replace(k, v)

    out, i = [], 0
    for m in _MATH_MARK.finditer(t):
        if m.start() > i:
            out.append((t[i:m.start()], "n"))
        if m.group(1):
            out.append((m.group(2), "sup" if m.group(1) == "^" else "sub"))
        else:
            out.append((m.group(4), "sup" if m.group(3) == "^" else "sub"))
        i = m.end()
    if i < len(t):
        out.append((t[i:], "n"))

    res = []
    for text, kind in out:
        text = text.replace("{", "").replace("}", "")
        text = re.sub(r"[ \t\u00a0]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        if text:
            res.append((text, kind))
    return res


def text_weight(s: str) -> float:
    """按「1 个汉字 = 1 em，1 个半角 ≈ 0.55 em」估算行宽。"""
    return sum(1.0 if ord(c) > 0x2E7F else 0.55 for c in s)


def fit_block(paras, w_in, h_in, candidates):
    """paras: [(text, size_ratio)]；返回能塞进 w×h 的最大基准字号（pt）。"""
    for s in candidates:
        need = 0.0
        for text, ratio in paras:
            if not text:
                continue
            fs = s * ratio
            per = max(1.0, (w_in * 72.0) / fs)
            n = sum(max(1, math.ceil(text_weight(ln) / per)) for ln in text.split("\n"))
            need += n * fs * 1.44
        if need <= h_in * 72.0:
            return s
    return candidates[-1]


# ────────────────────────────── 图片下载 / 压缩 ──────────────────────────────
_OPENER = urllib.request.build_opener(urllib.request.ProxyHandler({}))


def trim_white_border(im, thresh=30, pad=10, min_keep=0.18):
    """
    裁掉配图四周的白边（练习册那种"宽条裁片"大量空白，投屏时图会显得很小）。
    保守策略：识别到的内容框若小于原图 18%、或大于原图 92%，都放弃裁剪。
    """
    try:
        from PIL import Image, ImageChops
        g = im.convert("L")
        bg = Image.new("L", im.size, 255)
        mask = ImageChops.difference(g, bg).point(lambda v: 255 if v > thresh else 0)
        box = mask.getbbox()
        if not box:
            return im
        x0, y0, x1, y1 = box
        x0 = max(0, x0 - pad); y0 = max(0, y0 - pad)
        x1 = min(im.width, x1 + pad); y1 = min(im.height, y1 + pad)
        keep = ((x1 - x0) * (y1 - y0)) / float(im.width * im.height)
        if keep < min_keep or keep > 0.92:
            return im
        return im.crop((x0, y0, x1, y1))
    except Exception:
        return im


def fetch_image(url: str, cache: Path, max_px: int = 1800, quality: int = 82):
    if not url:
        return None
    name = re.sub(r"[^A-Za-z0-9._-]", "_", url.split("/")[-1]) or "img"
    name = f"{abs(hash(url)) % (10 ** 10)}_{name}"
    raw = cache / ("raw_" + name)
    out = cache / (Path(name).stem + ".jpg" if True else name)
    if out.exists() and out.stat().st_size > 0:
        return out
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "MinxueHandout/1.0"})
        with _OPENER.open(req, timeout=45) as r:
            data = r.read()
        try:
            from PIL import Image
            im = Image.open(io.BytesIO(data))
            if im.mode in ("RGBA", "LA", "P"):
                bg = Image.new("RGB", im.size, (255, 255, 255))
                im = im.convert("RGBA")
                bg.paste(im, mask=im.split()[-1])
                im = bg
            elif im.mode != "RGB":
                im = im.convert("RGB")
            im = trim_white_border(im)
            if max(im.size) > max_px:
                k = max_px / max(im.size)
                im = im.resize((max(1, int(im.width * k)), max(1, int(im.height * k))), Image.LANCZOS)
            im.save(out, "JPEG", quality=quality, optimize=True)
        except Exception:
            out.write_bytes(data)
        return out if out.exists() else None
    except Exception as e:
        print(f"    ! 图片下载失败 {e} :: {url[:80]}")
        return None


# ────────────────────────────── pptx 工具函数 ──────────────────────────────
from pptx import Presentation                                            # noqa: E402
from pptx.dml.color import RGBColor                                       # noqa: E402
from pptx.enum.shapes import MSO_SHAPE                                    # noqa: E402
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR                           # noqa: E402
from pptx.oxml.ns import qn, nsdecls                                      # noqa: E402
from pptx.oxml import parse_xml                                           # noqa: E402
from pptx.util import Inches, Pt, Emu                                     # noqa: E402


def rgb(t):
    return RGBColor(*t)


def style_run(run, size=None, bold=None, color=None, italic=None, font=FONT):
    f = run.font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = parse_xml(f'<{tag} {nsdecls("a")} typeface="{font}"/>')
            rPr.append(el)
        else:
            el.set("typeface", font)


def set_autofit(tf, mode="norm"):
    """统一 bodyPr 的自动调整：默认 normAutofit（文字超框时自动缩，而不是溢出压到答案框）。"""
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    for tag in ("a:normAutofit", "a:spAutoFit", "a:noAutofit"):
        el = bodyPr.find(qn(tag))
        if el is not None:
            bodyPr.remove(el)
    if mode:
        bodyPr.append(parse_xml(f'<a:{mode}Autofit {nsdecls("a")}/>'))


def add_math_runs(p, segs, size, color, bold=False):
    """按 segments 建多个 run；上下标用 a:rPr/@baseline 真排版（不用 Unicode 上下标字符）。"""
    for text, kind in segs:
        if not text:
            continue
        r = p.add_run()
        r.text = text
        if kind == "n":
            style_run(r, size=size, color=color, bold=bold)
        else:
            style_run(r, size=size * 0.62, color=color, bold=bold)
            r.font._rPr.set("baseline", "30000" if kind == "sup" else "-25000")


def add_box(slide, x, y, w, h, *, wrap=True, anchor=MSO_ANCHOR.TOP, margins=(0, 0, 0, 0), autofit="norm"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_right, tf.margin_top, tf.margin_bottom = [Inches(v) for v in margins]
    set_autofit(tf, autofit)
    return tb, tf


def para(tf, first):
    return tf.paragraphs[0] if first else tf.add_paragraph()


def fill_shape(shape, color, line_color=None, line_w=0.75):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line_color)
        shape.line.width = Pt(line_w)


def no_shadow(shape):
    spPr = shape._element.spPr
    el = parse_xml(f'<a:effectLst {nsdecls("a")}/>')
    spPr.append(el)


def fill_gradient(shape, c1, c2, angle=45):
    """双色线性渐变；失败时退回 c1 纯色。"""
    try:
        f = shape.fill
        f.gradient()
        stops = f.gradient_stops
        stops[0].position = 0.0
        stops[0].color.rgb = rgb(c1)
        stops[1].position = 1.0
        stops[1].color.rgb = rgb(c2)
        try:
            f.gradient_angle = angle
        except Exception:
            pass
    except Exception:
        fill_shape(shape, c1)


def set_fill_alpha(shape, pct):
    """给纯色填充加透明度，pct=0~100（不透明度）。"""
    try:
        spPr = shape._element.spPr
        sf = spPr.find(qn("a:solidFill"))
        if sf is None:
            return
        clr = sf.find(qn("a:srgbClr"))
        if clr is None:
            return
        for a in clr.findall(qn("a:alpha")):
            clr.remove(a)
        clr.append(parse_xml(f'<a:alpha {nsdecls("a")} val="{int(pct * 1000)}"/>'))
    except Exception:
        pass


def add_dots(slide, x, y, filled, color, n=4, d=0.105, gap=0.075):
    """难度圆点刻度。返回结束 x。"""
    for i in range(n):
        sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + i * (d + gap)), Inches(y), Inches(d), Inches(d))
        fill_shape(sp, color if i < (filled or 0) else DOT_OFF)
        no_shadow(sp)
    return x + n * (d + gap) - gap


def add_chip(slide, x, y, text, *, size=13, color=INK2, bg=WHITE, border=LINE, bold=False, pad=0.13, h=0.34):
    w = text_weight(text) * size / 72.0 + pad * 2 + 0.06
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.adjustments[0] = 0.5
    fill_shape(sp, bg, border, 0.75)
    no_shadow(sp)
    tf = sp.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    style_run(r, size=size, bold=bold, color=color)
    return sp, x + w + 0.10


def add_line(slide, x, y, w, color=LINE, h_pt=1.0):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Pt(h_pt))
    fill_shape(sp, color)
    no_shadow(sp)
    return sp


# ────────────────────────────── 单击浮现动画 ──────────────────────────────
def add_click_appear(slide, shape):
    """给 shape 注入「单击时淡入」动画（ECMA-376 p:timing）。"""
    spid = str(shape.shape_id)
    xml = f'''<p:timing {nsdecls("p", "a")}>
  <p:tnLst>
    <p:par>
      <p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">
        <p:childTnLst>
          <p:seq concurrent="1" nextAc="seek">
            <p:cTn id="2" dur="indefinite" nodeType="mainSeq">
              <p:childTnLst>
                <p:par>
                  <p:cTn id="3" fill="hold">
                    <p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>
                    <p:childTnLst>
                      <p:par>
                        <p:cTn id="4" fill="hold">
                          <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                          <p:childTnLst>
                            <p:par>
                              <p:cTn id="5" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="clickEffect">
                                <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                                <p:childTnLst>
                                  <p:set>
                                    <p:cBhvr>
                                      <p:cTn id="6" dur="1" fill="hold">
                                        <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                                      </p:cTn>
                                      <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
                                      <p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>
                                    </p:cBhvr>
                                    <p:to><p:strVal val="visible"/></p:to>
                                  </p:set>
                                  <p:animEffect transition="in" filter="fade">
                                    <p:cBhvr>
                                      <p:cTn id="7" dur="400"/>
                                      <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
                                    </p:cBhvr>
                                  </p:animEffect>
                                </p:childTnLst>
                              </p:cTn>
                            </p:par>
                          </p:childTnLst>
                        </p:cTn>
                      </p:par>
                    </p:childTnLst>
                  </p:cTn>
                </p:par>
              </p:childTnLst>
            </p:cTn>
            <p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
            <p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>
          </p:seq>
        </p:childTnLst>
      </p:cTn>
    </p:par>
  </p:tnLst>
</p:timing>'''
    el = parse_xml(xml)
    slide._element.append(el)


# ────────────────────────────── 版面：封面 / 目录 / 分节 / 题目 ──────────────────────────────
def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def bg_rect(slide, color=WHITE):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SLIDE_W), Inches(SLIDE_H))
    fill_shape(sp, color)
    no_shadow(sp)
    return sp


def slide_cover(prs, data):
    s = blank(prs)
    bg = bg_rect(s, GRAD1)
    fill_gradient(bg, GRAD1, GRAD2, 35)
    # 装饰圆
    for cx, cy, d, a in ((10.0, -1.9, 5.4, 9), (11.8, 4.6, 3.8, 7), (-1.5, 5.2, 3.2, 6), (8.9, 5.9, 1.6, 8)):
        c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx), Inches(cy), Inches(d), Inches(d))
        fill_shape(c, WHITE)
        no_shadow(c)
        set_fill_alpha(c, a)

    tb, tf = add_box(s, 0.92, 0.86, 11.5, 2.1)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = "周末错题讲评"
    style_run(r, size=20, bold=True, color=(0xCF, 0xCB, 0xF7))
    p1 = tf.add_paragraph(); p1.space_before = Pt(6)
    r1 = p1.add_run(); r1.text = data.get("title", "周末班错题课件")
    style_run(r1, size=44, bold=True, color=WHITE)
    p2 = tf.add_paragraph(); p2.space_before = Pt(10)
    r2 = p2.add_run(); r2.text = "日期倒序 · 每天内部由易到难　｜　放映时单击一下浮现参考答案"
    style_run(r2, size=15, color=(0xC9, 0xC6, 0xF2))

    st = data.get("stats", {})
    nums = [
        ("题目", str(st.get("topics", 0))),
        ("有课日", str(st.get("sections", 0))),
        ("学生", str(st.get("students", 0))),
        ("错题条数", str(st.get("rawRows", 0))),
    ]
    x = 0.92
    for label, val in nums:
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(3.42), Inches(2.52), Inches(1.42))
        card.adjustments[0] = 0.14
        fill_shape(card, WHITE)
        no_shadow(card)
        set_fill_alpha(card, 14)
        tfc = card.text_frame
        tfc.margin_left = tfc.margin_right = Inches(0.22)
        tfc.margin_top = Inches(0.18); tfc.margin_bottom = Inches(0.14)
        pc = tfc.paragraphs[0]
        rc = pc.add_run(); rc.text = val
        style_run(rc, size=31, bold=True, color=WHITE)
        pl = tfc.add_paragraph()
        rl = pl.add_run(); rl.text = label
        style_run(rl, size=12.5, color=(0xD4, 0xD1, 0xF6))
        x += 2.78

    per = data.get("period", {})
    meta = [
        f"时段　{per.get('start', '')} ~ {per.get('end', '')}",
        f"学生　{'、'.join(st.get('studentNames', []))}",
        ("备注　演讲者视图可见：谁错了、错在哪、原卷图链接" ),
    ]
    tb3, tf3 = add_box(s, 0.92, 5.30, 11.4, 1.7)
    first = True
    for ln in meta:
        p = para(tf3, first); first = False
        p.space_after = Pt(8)
        r = p.add_run(); r.text = ln
        style_run(r, size=14, color=(0xBE, 0xBA, 0xEE))
    return s


def slide_overview(prs, data):
    s = blank(prs)
    bg_rect(s, PAGE_BG)
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(M), Inches(0.42), Inches(SLIDE_W - 2 * M), Inches(SLIDE_H - 0.84))
    card.adjustments[0] = 0.030
    fill_shape(card, WHITE, CARD_LINE, 1.0)
    no_shadow(card)

    bar = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(M + 0.42), Inches(0.86), Inches(0.085), Inches(0.46))
    bar.adjustments[0] = 0.5
    fill_gradient(bar, GRAD1, GRAD2, 90)
    no_shadow(bar)
    tb, tf = add_box(s, M + 0.66, 0.80, 10.5, 0.62)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = "本份课件目录"
    style_run(r, size=25, bold=True, color=INK)
    r0 = p.add_run(); r0.text = "　按日期倒序"
    style_run(r0, size=14, color=INK3)

    rows = data.get("overview", [])
    tb2, tf2 = add_box(s, M + 0.44, 1.62, SLIDE_W - 2 * M - 0.88, 4.9)
    head = f"{'日期':<20}{'题数':>5}{'学生':>6}{'基础':>6}{'中等':>6}{'较难':>6}{'未判定':>7}"
    p = tf2.paragraphs[0]
    r = p.add_run(); r.text = head
    style_run(r, size=13, bold=True, color=ACCENT)
    add_line(s, M + 0.44, 2.12, SLIDE_W - 2 * M - 0.88, CARD_LINE, 1.2)
    for i, row in enumerate(rows):
        p = tf2.add_paragraph()
        p.space_after = Pt(4)
        r = p.add_run()
        r.text = (f"{row.get('label',''):<20}{row.get('topics',0):>5}{row.get('students',0):>6}"
                  f"{row.get('basic',0):>6}{row.get('medium',0):>6}{row.get('hard',0):>6}{row.get('unknown',0):>7}")
        style_run(r, size=14, color=INK if i % 2 == 0 else INK2)
    return s


def slide_section(prs, sec, seq_no, total_sec):
    s = blank(prs)
    bg = bg_rect(s, GRAD1)
    fill_gradient(bg, GRAD1, GRAD2, 35)
    for cx, cy, d, a in ((10.6, -1.4, 4.6, 9), (-1.2, 5.0, 3.4, 7)):
        c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx), Inches(cy), Inches(d), Inches(d))
        fill_shape(c, WHITE)
        no_shadow(c)
        set_fill_alpha(c, a)

    tb, tf = add_box(s, M + 0.46, 2.14, 11.3, 3.0)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = f"SECTION {seq_no:02d} / {total_sec:02d}"
    style_run(r, size=15, bold=True, color=(0xCF, 0xCB, 0xF7))
    p1 = tf.add_paragraph(); p1.space_before = Pt(8)
    r1 = p1.add_run(); r1.text = sec.get("label", "")
    style_run(r1, size=46, bold=True, color=WHITE)
    p2 = tf.add_paragraph(); p2.space_before = Pt(18)
    r2 = p2.add_run()
    tiers = sec.get("tiers", {})
    tb2 = [f"{lab} {tiers[k]}" for k, lab in (("basic", "基础"), ("medium", "中等"), ("hard", "较难"), ("unknown", "未判定")) if tiers.get(k)]
    r2.text = f"{sec.get('topicCount',0)} 题" + (f"　·　{' · '.join(tb2)}" if tb2 else "") + f"　·　{sec.get('studentCount',0)} 名学生"
    style_run(r2, size=16, color=(0xD4, 0xD1, 0xF6))
    p3 = tf.add_paragraph(); p3.space_before = Pt(8)
    r3 = p3.add_run(); r3.text = "错的学生：" + "、".join(sec.get("students", []))
    style_run(r3, size=14, color=(0xBE, 0xBA, 0xEE))
    return s


def slide_question(prs, q, fig_path, page_img_path, animate=True, with_answer=True):
    s = blank(prs)
    bg_rect(s, PAGE_BG)

    tier = q.get("tier", "unknown")
    tcol = TIER_COLOR.get(tier, TIER_COLOR["unknown"])

    # ── 顶部：编号徽章 + 两行元信息（右侧学生名单） ──
    badge = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(M), Inches(0.26), Inches(0.64), Inches(0.64))
    badge.adjustments[0] = 0.30
    fill_gradient(badge, GRAD1, GRAD2, 55)
    no_shadow(badge)
    tfi = badge.text_frame
    tfi.margin_left = tfi.margin_right = tfi.margin_top = tfi.margin_bottom = 0
    tfi.vertical_anchor = MSO_ANCHOR.MIDDLE
    pi = tfi.paragraphs[0]
    pi.alignment = PP_ALIGN.CENTER
    ri = pi.add_run()
    ri.text = str(q.get("index", 0))
    style_run(ri, size=22, bold=True, color=WHITE)

    hx = M + 0.84
    tbm, tfm = add_box(s, hx, 0.27, 7.6, 0.30, wrap=False, autofit=None)
    pm = tfm.paragraphs[0]
    segs1 = [q.get("dayLabel", "")]
    if q.get("typeLabel"):
        segs1.append(q["typeLabel"])
    rm = pm.add_run()
    rm.text = "　·　".join(x for x in segs1 if x)
    style_run(rm, size=12.5, color=INK3)

    diff = q.get("difficulty")
    x2 = hx
    y2 = 0.575
    if diff is None:
        _, x2 = add_chip(s, x2, y2 - 0.02, "难度未判定", size=11.5, color=WARN_INK, bg=WARN_BG, border=None, h=0.30)
    else:
        xend = add_dots(s, x2, y2 + 0.06, diff, tcol)
        tbdd, tfdd = add_box(s, xend + 0.10, y2 - 0.04, 2.2, 0.32, wrap=False, autofit=None)
        pdd = tfdd.paragraphs[0]
        rdd = pdd.add_run()
        rdd.text = f"难度 {diff} · {q.get('tierLabel', '')}"
        style_run(rdd, size=12.5, bold=True, color=tcol)
        x2 = xend + 0.10 + text_weight(rdd.text) * 12.5 / 72.0 + 0.30
    n = q.get("studentCount", 1)
    if n >= 2:
        _, x2 = add_chip(s, x2, y2 - 0.02, f"{n} 人错", size=11.5, bold=True,
                         color=(0xB9, 0x1C, 0x1C), bg=(0xFC, 0xE9, 0xE9), border=None, h=0.30)
    if q.get("diffInconsistent"):
        dv = "、".join(str(v) for v in (q.get("diffValues") or []))
        _, x2 = add_chip(s, x2, y2 - 0.02, f"标注不一致 {dv}", size=11,
                         color=(0x6D, 0x28, 0xD9), bg=(0xF5, 0xF3, 0xFF), border=None, h=0.30)

    names = "、".join(st.get("name", "") for st in q.get("students", []))
    if q.get("questionNumber") is not None:
        names = f"卷面第 {q['questionNumber']} 题" + (f"　·　{names}" if names else "")
    tbw, tfw = add_box(s, SLIDE_W - M - 5.4, 0.38, 5.4, 0.46, wrap=False, autofit=None)
    pw = tfw.paragraphs[0]
    pw.alignment = PP_ALIGN.RIGHT
    rw = pw.add_run()
    rw.text = names
    style_run(rw, size=13.5, bold=True, color=INK2)

    # ── 内容卡片 ──
    card_y = 1.06
    card_bottom = ANS_Y - 0.14
    card_h = card_bottom - card_y
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(M), Inches(card_y), Inches(SLIDE_W - 2 * M), Inches(card_h))
    card.adjustments[0] = 0.030
    fill_shape(card, WHITE, CARD_LINE, 1.0)
    no_shadow(card)

    ix = M + 0.32
    iw_total = SLIDE_W - 2 * M - 0.64

    has_fig = bool(fig_path)
    has_page = bool(page_img_path)
    fig_hint = bool(FIG_HINT.search((q.get("parentStem") or "") + (q.get("stem") or "")))
    # 配图栏宽度随图片长宽比自适应（竖长条收窄、横图放宽），让配图尽量占满卡片高度
    panel_w = 0.0
    if has_fig or has_page:
        img0 = fig_path or page_img_path
        try:
            from PIL import Image as _Img
            iw0, ih0 = _Img.open(img0).size
            ar = ih0 / max(1, iw0)
        except Exception:
            ar = 0.72
        avail_h0 = (card_h - 0.44) - 0.62
        want = avail_h0 / max(ar, 0.05) + 0.36
        panel_w = max(2.7, min(6.7, want, iw_total - 4.6))
    text_w = iw_total - (panel_w + 0.30 if panel_w else 0)

    ycur = card_y + 0.24

    # 共享题干缺失警示（多小问拆分行，公共条件未入库）
    stem_missing = bool(q.get("subNo")) and not q.get("parentStem")

    cond_h = 0.0
    if q.get("parentStem"):
        cond_plain = norm_math(q["parentStem"])
        per_line = max(1.0, (text_w - 0.66) * 72.0 / 14.0)
        lines = sum(max(1, math.ceil(text_weight(ln) / per_line)) for ln in cond_plain.split("\n"))
        cond_h = min(1.9, lines * 14 * 1.5 / 72.0 + 0.34)
        cond_box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(ix), Inches(ycur), Inches(text_w), Inches(cond_h))
        cond_box.adjustments[0] = min(0.5, 0.10 / max(cond_h, 0.1))
        fill_shape(cond_box, COND_BG)
        no_shadow(cond_box)
        bar = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(ix), Inches(ycur), Inches(0.065), Inches(cond_h))
        bar.adjustments[0] = 0.5
        fill_shape(bar, COND_BAR)
        no_shadow(bar)
        tfcd = cond_box.text_frame
        tfcd.margin_left = Inches(0.24)
        tfcd.margin_right = Inches(0.16)
        tfcd.margin_top = Inches(0.07)
        tfcd.margin_bottom = Inches(0.07)
        tfcd.vertical_anchor = MSO_ANCHOR.MIDDLE
        set_autofit(tfcd, "norm")
        pl0 = tfcd.paragraphs[0]
        rl0 = pl0.add_run()
        rl0.text = "公共题干"
        style_run(rl0, size=10.5, bold=True, color=COND_INK)
        pl1 = tfcd.add_paragraph()
        add_math_runs(pl1, math_segments(q["parentStem"]), 13.5, COND_INK)
        ycur += cond_h + 0.18

    if stem_missing:
        wbox = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(ix), Inches(ycur), Inches(text_w), Inches(0.44))
        wbox.adjustments[0] = 0.22
        fill_shape(wbox, WARN_BG, WARN_LINE, 0.75)
        no_shadow(wbox)
        tfw2 = wbox.text_frame
        tfw2.margin_left = Inches(0.18)
        tfw2.margin_right = Inches(0.12)
        tfw2.margin_top = tfw2.margin_bottom = 0
        tfw2.vertical_anchor = MSO_ANCHOR.MIDDLE
        pw2 = tfw2.paragraphs[0]
        rw2 = pw2.add_run()
        rw2.text = "⚠ 该题是多小问拆分行，共享题干未入库 —— 请对照原卷图先补条件，再讲本题"
        style_run(rw2, size=12.5, bold=True, color=WARN_INK)
        ycur += 0.44 + 0.16

    # ── 题干 + 选项 ──
    stem = q.get("stem") or ""
    if not norm_math(stem):
        stem = "（题干为空，请对照原卷图）"
    paras = []
    if q.get("stemIsFallback"):
        paras.append((math_segments("⚠ 原题干未识别到，以下文字来自作答内容，请对照原卷图"), 0.60))
    paras.append((math_segments(stem), 1.0))
    for i, o in enumerate(q.get("options") or []):
        paras.append((math_segments(f"{'ABCDE'[i]}. {o}"), 0.85))

    body_top = ycur
    body_h = card_bottom - 0.24 - body_top
    flat = [("".join(t for t, _ in segs), ratio) for segs, ratio in paras]
    base = fit_block(flat, text_w - 0.1, body_h, [30, 28, 26, 24, 22, 20, 18, 16, 15, 14, 13, 12])

    tbb, tfb = add_box(s, ix, body_top, text_w, body_h)
    first = True
    for segs, ratio in paras:
        if not segs:
            continue
        p = para(tfb, first)
        first = False
        sz = base * ratio
        p.space_after = Pt(max(1.2, sz * (0.20 if ratio == 1.0 else 0.13)))
        add_math_runs(p, segs, sz, INK3 if ratio < 0.8 else INK)

    # ── 题图（只贴系统几何裁片；默认**不贴整页原卷**）──
    if panel_w:
        px = M + (SLIDE_W - 2 * M) - 0.26 - panel_w
        py = card_y + 0.22
        ph = card_h - 0.44
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(px), Inches(py), Inches(panel_w), Inches(ph))
        box.adjustments[0] = 0.035
        fill_shape(box, PAGE_BG, CARD_LINE, 0.75)
        no_shadow(box)
        cap = "题图（系统裁片）"
        img = fig_path or page_img_path
        from PIL import Image
        try:
            iw, ih = Image.open(img).size
        except Exception:
            iw, ih = 4, 3
        avail_w = panel_w - 0.34
        avail_h = ph - 0.62
        k = min(avail_w / iw, avail_h / ih)
        w_in, h_in = iw * k, ih * k
        s.shapes.add_picture(str(img), Inches(px + (panel_w - w_in) / 2), Inches(py + 0.18), Inches(w_in), Inches(h_in))
        tbc, tfc = add_box(s, px + 0.16, py + ph - 0.40, panel_w - 0.32, 0.30)
        pc = tfc.paragraphs[0]
        pc.alignment = PP_ALIGN.CENTER
        rc = pc.add_run()
        rc.text = cap
        style_run(rc, size=10.5, color=INK3)
    elif q.get("figHint"):
        # 原卷有图但系统没裁出配图：不贴原卷页，只给老师一句提示（原卷链接在备注页）
        tbn, tfn = add_box(s, ix, card_bottom - 0.62, text_w, 0.36, wrap=False, autofit=None)
        pn = tfn.paragraphs[0]
        rn = pn.add_run()
        rn.text = "本题原卷含图形，系统未裁出配图 —— 讲前请自备图形（原卷链接见备注页）"
        style_run(rn, size=11.5, color=WARN_INK)

    # ── 参考答案（单击浮现） ──
    ans_box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(M), Inches(ANS_Y), Inches(SLIDE_W - 2 * M), Inches(ANS_H))
    ans_box.adjustments[0] = 0.14
    if not with_answer or not q.get("hasAnswer"):
        missing = with_answer
        fill_shape(ans_box, WARN_BG if missing else WHITE, WARN_LINE if missing else CARD_LINE, 1.0)
    else:
        fill_shape(ans_box, (0xEB, 0xFB, 0xF3), (0xBF, 0xE9, 0xD6), 1.0)
    no_shadow(ans_box)
    tfa = ans_box.text_frame
    tfa.word_wrap = True
    tfa.margin_left = tfa.margin_right = Inches(0.26)
    tfa.margin_top = tfa.margin_bottom = Inches(0.09)
    tfa.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_autofit(tfa, "norm")

    if not with_answer:
        pa = tfa.paragraphs[0]
        ra = pa.add_run()
        ra.text = "（本份为无答案版）"
        style_run(ra, size=15, color=INK3)
    elif q.get("hasAnswer"):
        ans_text = norm_math(q.get("answer") or "")
        label = "参考答案"
        src = q.get("answerSourceLabel") or ""
        if src:
            label += f"　·　{src}"
        if q.get("answerRisk"):
            label += f"　·　⚠ {q['answerRisk']}"
        if len(ans_text) > 160:
            label += "　·　⚠ 答案较长，讲前请核对"
        pa = tfa.paragraphs[0]
        ra = pa.add_run()
        ra.text = label
        style_run(ra, size=11, bold=True, color=ANS_INK if not q.get("answerRisk") else WARN_INK)
        sz = fit_block([(ans_text, 1.0), (label, 0.6)], SLIDE_W - 2 * M - 0.6, ANS_H - 0.22,
                       [26, 24, 22, 20, 18, 16, 15, 14, 13, 12, 11])
        pa2 = tfa.add_paragraph()
        pa2.space_before = Pt(3)
        add_math_runs(pa2, math_segments(q.get("answer") or ""), sz, ANS_INK, bold=True)
    else:
        pa = tfa.paragraphs[0]
        ra = pa.add_run()
        ra.text = "参考答案"
        style_run(ra, size=11, bold=True, color=WARN_INK)
        pa2 = tfa.add_paragraph()
        pa2.space_before = Pt(3)
        ra2 = pa2.add_run()
        ra2.text = "库里为空 —— 讲前请人工补"
        style_run(ra2, size=18, bold=True, color=WARN_INK)

    if animate and with_answer:
        try:
            add_click_appear(s, ans_box)
        except Exception as e:
            print(f"    ! 动画注入失败（{e}），该页答案将直接可见")

    # ── 演讲者备注 ──
    note = [f"第 {q.get('index',0)} 题 · {q.get('dayLabel','')} · "
            f"{'难度 ' + str(q.get('difficulty')) if q.get('difficulty') is not None else '难度未判定'} · "
            f"{q.get('studentCount',1)} 人错 · {q.get('typeLabel','')}"]
    if q.get("parentStem"):
        note.append("母题：" + norm_math(q["parentStem"]))
    for st in q.get("students", []):
        bits = [f"· {st.get('name','')}"]
        if st.get("wrongTimes", 1) > 1:
            bits.append(f"（错 {st['wrongTimes']} 次）")
        if st.get("blank"):
            bits.append("未作答")
        elif st.get("answer"):
            bits.append("答：" + norm_math(str(st["answer"])))
        if st.get("errorType"):
            bits.append(f"[{st['errorType']}]")
        if st.get("errorReason"):
            bits.append(st["errorReason"])
        note.append(" ".join(bits))
    if q.get("analysis"):
        note.append("解析：" + norm_math(q["analysis"])[:400])
    doc = next((st.get("docImage") for st in q.get("students", []) if st.get("docImage")), None)
    if doc:
        note.append("原卷图：" + doc)
    s.notes_slide.notes_text_frame.text = "\n".join(note)
    return s


# ────────────────────────────── PDF（HTML + 无头浏览器） ──────────────────────────────
def build_print_html(data, slides, imgmap, out_html: Path, with_answer: bool):
    def e(x):
        return (str(x or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))

    def hmath(s):
        """题干/答案的富文本：上下标用 <sup>/<sub> 真排版。"""
        out = []
        for text, kind in math_segments(s):
            t = e(text)
            if kind == "sup":
                out.append(f"<sup>{t}</sup>")
            elif kind == "sub":
                out.append(f"<sub>{t}</sub>")
            else:
                out.append(t)
        return "".join(out)

    def uri(p):
        return Path(p).as_uri() if p else ""

    body = []
    st = data.get("stats", {})
    per = data.get("period", {})
    body.append(f'''<section class="slide cover">
  <div class="cv-circ c1"></div><div class="cv-circ c2"></div><div class="cv-circ c3"></div>
  <div class="cv-inner">
    <p class="cv-kicker">周末错题讲评</p>
    <h1>{e(data.get("title"))}</h1>
    <p class="cv-sub">日期倒序 · 每天内部由易到难　｜　PDF 版答案直接印出</p>
    <div class="cv-nums">{''.join(f'<div class="cv-card"><b>{v}</b><span>{k}</span></div>' for k, v in
        [("题目", st.get("topics",0)), ("有课日", st.get("sections",0)), ("学生", st.get("students",0)), ("错题条数", st.get("rawRows",0))])}</div>
    <ul class="cv-meta">
      <li>时段　{e(per.get("start"))} ~ {e(per.get("end"))}</li>
      <li>学生　{e("、".join(st.get("studentNames", [])))}</li>
      <li>备注　演讲者视图可见：谁错了、错在哪、原卷图链接</li>
    </ul>
  </div>
</section>''')

    rows = "".join(
        f'<tr><td>{e(r.get("label"))}</td><td>{r.get("topics",0)}</td><td>{r.get("students",0)}</td>'
        f'<td>{r.get("basic",0)}</td><td>{r.get("medium",0)}</td><td>{r.get("hard",0)}</td><td>{r.get("unknown",0)}</td></tr>'
        for r in data.get("overview", []))
    body.append(f'''<section class="slide overview">
  <div class="ocard">
    <h2><span class="bar"></span>本份课件目录<small>按日期倒序</small></h2>
    <table><thead><tr><th>日期</th><th>题数</th><th>学生</th><th>基础</th><th>中等</th><th>较难</th><th>未判定</th></tr></thead>
    <tbody>{rows}</tbody></table>
  </div>
</section>''')

    sec_no = 0
    for sl in slides:
        if sl["kind"] == "section":
            sec_no += 1
            tiers = sl.get("tiers", {})
            tb = [f"{lab} {tiers[k]}" for k, lab in (("basic", "基础"), ("medium", "中等"), ("hard", "较难"), ("unknown", "未判定")) if tiers.get(k)]
            bits = f'{sl.get("topicCount",0)} 题' + (f'　·　{" · ".join(tb)}' if tb else "") + f'　·　{sl.get("studentCount",0)} 名学生'
            body.append(f'''<section class="slide divider">
  <div class="dv-circ c1"></div><div class="dv-circ c2"></div>
  <p class="kicker">SECTION {sec_no:02d} / {len([x for x in slides if x["kind"]=="section"]):02d}</p>
  <h1>{e(sl.get("label"))}</h1>
  <p class="dim">{e(bits)}</p>
  <p class="dim2">错的学生：{e("、".join(sl.get("students", [])))}</p>
</section>''')
            continue

        q = sl
        tier = q.get("tier", "unknown")
        diff = q.get("difficulty")
        n = q.get("studentCount", 1)
        dots = "".join(f'<i class="{"on" if i < (diff or 0) else ""}"></i>' for i in range(4))
        if diff is None:
            diff_html = '<span class="chip warn">难度未判定</span>'
        else:
            diff_html = (f'<span class="dots d-{tier}">{dots}</span>'
                         f'<b class="tl t-{tier}">难度 {diff} · {e(q.get("tierLabel"))}</b>')
        hot = f'<span class="hot">{n} 人错</span>' if n >= 2 else f'<span class="dimc">{n} 人错</span>'
        jit = (f'<span class="chip jit">标注不一致 {e("、".join(str(v) for v in (q.get("diffValues") or [])))}</span>'
               if q.get("diffInconsistent") else "")
        names = "、".join(x.get("name", "") for x in q.get("students", []))
        if q.get("questionNumber") is not None:
            names = f'卷面第 {q["questionNumber"]} 题' + (f'　·　{names}' if names else "")

        stem_missing = bool(q.get("subNo")) and not q.get("parentStem")
        warn_html = ('<div class="miss-warn">⚠ 该题是多小问拆分行，共享题干未入库 —— 请对照原卷图先补条件，再讲本题</div>'
                     if stem_missing else "")

        parts = []
        paras = []
        if q.get("parentStem"):
            paras.append((norm_math(q["parentStem"]), 0.72))
        raw_stem = q.get("stem") or ""
        if not norm_math(raw_stem):
            raw_stem = "（题干为空，请对照原卷图）"
        if q.get("stemIsFallback"):
            paras.append(("⚠ 原题干未识别到，以下文字来自作答内容，请对照原卷图", 0.62))
        paras.append((norm_math(raw_stem), 1.0))
        for i, o in enumerate(q.get("options") or []):
            paras.append((norm_math(o), 0.85))

        figp = imgmap.get(id(q), (None, None))
        fig_html = ""
        nofig_html = ""
        panel_w = 0.0
        if figp[0] or figp[1]:
            src = uri(figp[0] or figp[1])
            try:
                from PIL import Image as _Img
                iw0, ih0 = _Img.open(figp[0] or figp[1]).size
                ar = ih0 / max(1, iw0)
            except Exception:
                ar = 0.72
            avail_h0 = (ANS_Y - 0.14 - 1.06) - 0.44 - 0.62
            want = avail_h0 / max(ar, 0.05) + 0.36
            panel_w = max(2.7, min(6.7, want, (SLIDE_W - 2 * M - 0.64) - 4.6))
            fig_html = (f'<div class="panel" style="width:{panel_w:.2f}in"><img src="{src}"/>'
                        f'<div class="cap">题图（系统裁片）</div></div>')
        elif FIG_HINT.search((q.get("parentStem") or "") + (q.get("stem") or "")):
            nofig_html = '<div class="nofig">本题原卷含图形，系统未裁出配图 —— 讲前请自备图形（原卷链接见 PPT 备注页）</div>'

        text_w = SLIDE_W - 2 * M - ((panel_w + GUTTER) if panel_w else 0)
        base = fit_block(paras, text_w - 0.1, BODY_H - (0.9 if q.get("parentStem") else 0) - (0.5 if stem_missing else 0),
                         [30, 28, 26, 24, 22, 20, 18, 16, 15, 14, 13, 12])

        cond_html = f'<div class="cond"><span>公共题干</span>{hmath(q["parentStem"])}</div>' if q.get("parentStem") else ""
        fb_html = '<div class="fallback">⚠ 原题干未识别到，以下文字来自作答内容，请对照原卷图</div>' if q.get("stemIsFallback") else ""

        inner = [cond_html, warn_html, fb_html, f'<div class="stem">{hmath(raw_stem)}</div>']
        for i, o in enumerate(q.get("options") or []):
            inner.append(f'<div class="opt"><i>{"ABCDE"[i]}.</i> {hmath(o)}</div>')
        if nofig_html:
            inner.append(nofig_html)

        if not with_answer:
            ans_html = '<div class="ans plain">（本份为无答案版）</div>'
        elif q.get("hasAnswer"):
            lab = "参考答案"
            if q.get("answerSourceLabel"):
                lab += f'　·　{e(q["answerSourceLabel"])}'
            ans_plain = norm_math(q.get("answer"))
            if len(ans_plain) > 160:
                lab += '　·　⚠ 答案较长，讲前请核对'
            risk = f'<span class="risk">　·　⚠ {e(q["answerRisk"])}</span>' if q.get("answerRisk") else ""
            asz = fit_block([(ans_plain, 1.0), (lab, 0.6)], SLIDE_W - 2 * M - 0.6, ANS_H - 0.22,
                            [26, 24, 22, 20, 18, 16, 15, 14, 13, 12, 11])
            ans_html = (f'<div class="ans"><div class="lab">{lab}{risk}</div>'
                        f'<div class="val" style="font-size:{asz}pt">{hmath(q.get("answer"))}</div></div>')
        else:
            ans_html = '<div class="ans miss"><div class="lab">参考答案</div><div class="val">库里为空 —— 讲前请人工补</div></div>'

        m1 = e(q.get("dayLabel") or "")
        if q.get("typeLabel"):
            m1 += f' · {e(q["typeLabel"])}'
        content_html = "".join(inner)
        card_inner = (f'<div class="crow"><div class="ccol">{content_html}</div>{fig_html}</div>'
                      if fig_html else content_html)
        body.append(f'''<section class="slide q">
  <div class="hdr">
    <div class="idx">{q.get("index")}</div>
    <div class="meta"><div class="m1">{m1}</div><div class="m2">{diff_html}{hot}{jit}</div></div>
    <div class="who">{e(names)}</div>
  </div>
  <div class="card" style="font-size:{base}pt">{card_inner}</div>
  {ans_html}
</section>''')

    css = f'''
  @page {{ size: {SLIDE_W}in {SLIDE_H}in; margin: 0; }}
  * {{ box-sizing: border-box; -webkit-print-color-adjust: exact; print-color-adjust: exact; }}
  html,body {{ margin:0; padding:0; font-family:"{FONT}","Microsoft YaHei",sans-serif; color:#1b1f24; }}
  .slide {{ width:{SLIDE_W}in; height:{SLIDE_H}in; overflow:hidden; page-break-after:always; break-after:page;
            position:relative; background:#f2f4fb; padding:.30in {M}in .32in; display:flex; flex-direction:column; }}
  .slide:last-child {{ page-break-after:auto; break-after:auto; }}

  /* ── 封面 ── */
  .cover {{ padding:0; background:linear-gradient(135deg,#4a4ee4 0%,#7c38c9 100%); color:#fff; }}
  .cv-circ {{ position:absolute; border-radius:50%; background:rgba(255,255,255,.09); }}
  .cv-circ.c1 {{ width:5.4in; height:5.4in; right:-1.2in; top:-1.9in; }}
  .cv-circ.c2 {{ width:3.8in; height:3.8in; right:-0.6in; bottom:-1.2in; }}
  .cv-circ.c3 {{ width:3.2in; height:3.2in; left:-1.5in; bottom:-1.0in; background:rgba(255,255,255,.06); }}
  .cv-inner {{ position:relative; padding:.86in .92in; height:100%; display:flex; flex-direction:column; }}
  .cv-kicker {{ margin:0; font-size:19pt; font-weight:700; color:#cfcbf7; letter-spacing:.06em; }}
  .cover h1 {{ margin:.10in 0 .12in; font-size:44pt; line-height:1.15; }}
  .cv-sub {{ margin:0 0 .38in; font-size:15pt; color:#c9c6f2; }}
  .cv-nums {{ display:flex; gap:.26in; }}
  .cv-card {{ background:rgba(255,255,255,.13); border-radius:14px; padding:.16in .24in; width:2.52in; }}
  .cv-card b {{ display:block; font-size:31pt; line-height:1.1; }}
  .cv-card span {{ color:#d4d1f6; font-size:12.5pt; }}
  .cv-meta {{ margin:auto 0 0; padding:0; list-style:none; color:#bebaee; font-size:14pt; line-height:1.9; }}

  /* ── 目录 ── */
  .overview {{ background:#f2f4fb; }}
  .overview .ocard {{ background:#fff; border:1px solid #e5e9f4; border-radius:12px; flex:1;
                      padding:.40in .44in; display:flex; flex-direction:column; }}
  .overview h2 {{ margin:0 0 .18in; font-size:25pt; display:flex; align-items:center; gap:.14in; }}
  .overview h2 .bar {{ width:.085in; height:.46in; border-radius:99px; background:linear-gradient(180deg,#4a4ee4,#7c38c9); }}
  .overview h2 small {{ font-size:14pt; color:#8b949e; font-weight:400; }}
  .overview table {{ width:100%; border-collapse:collapse; font-size:14pt; }}
  .overview th,.overview td {{ padding:8px 10px; border-bottom:1px solid #eef1f7; text-align:right; }}
  .overview th:first-child,.overview td:first-child {{ text-align:left; }}
  .overview th {{ color:#4c51e0; font-size:12.5pt; font-weight:700; }}

  /* ── 分节页 ── */
  .divider {{ background:linear-gradient(135deg,#4a4ee4 0%,#7c38c9 100%); color:#fff; justify-content:center; padding-left:{M+0.46}in; }}
  .dv-circ {{ position:absolute; border-radius:50%; background:rgba(255,255,255,.08); }}
  .dv-circ.c1 {{ width:4.6in; height:4.6in; right:-1.0in; top:-1.4in; }}
  .dv-circ.c2 {{ width:3.4in; height:3.4in; left:-1.2in; bottom:-1.4in; background:rgba(255,255,255,.06); }}
  .divider .kicker {{ margin:0 0 .10in; font-size:15pt; font-weight:700; letter-spacing:.08em; color:#cfcbf7; }}
  .divider h1 {{ margin:0 0 .22in; font-size:46pt; }}
  .divider .dim {{ color:#d4d1f6; font-size:16pt; margin:0 0 .10in; }}
  .divider .dim2 {{ color:#bebaee; font-size:14pt; margin:0; }}

  /* ── 题目页 ── */
  .hdr {{ display:flex; align-items:center; gap:.18in; height:.64in; margin-bottom:.14in; }}
  .idx {{ width:.64in; height:.64in; border-radius:12px; flex:none;
          background:linear-gradient(135deg,#4a4ee4,#7c38c9); color:#fff; font-size:22pt; font-weight:700;
          display:flex; align-items:center; justify-content:center; }}
  .meta .m1 {{ font-size:12.5pt; color:#8b949e; line-height:1.3; }}
  .meta .m2 {{ display:flex; align-items:center; gap:.14in; margin-top:.035in; }}
  .dots {{ display:inline-flex; gap:.055in; }}
  .dots i {{ width:.105in; height:.105in; border-radius:50%; background:#dde2f0; }}
  .dots.d-basic i.on {{ background:#0f7b4f; }}
  .dots.d-medium i.on {{ background:#1e6fd9; }}
  .dots.d-hard i.on {{ background:#c2410c; }}
  .dots.d-unknown i.on {{ background:#6b7280; }}
  .tl {{ font-size:12.5pt; }}
  .t-basic {{ color:#0f7b4f; }} .t-medium {{ color:#1e6fd9; }} .t-hard {{ color:#c2410c; }} .t-unknown {{ color:#6b7280; }}
  .hot {{ font-size:11.5pt; font-weight:700; color:#b91c1c; background:#fce9e9; border-radius:999px; padding:.02in .12in; }}
  .dimc {{ font-size:11.5pt; color:#8b949e; }}
  .chip {{ font-size:11pt; padding:.02in .11in; border-radius:999px; white-space:nowrap; }}
  .chip.warn {{ color:#b45309; background:#fffbeb; }}
  .chip.jit {{ color:#6d28d9; background:#f5f3ff; }}
  .who {{ margin-left:auto; font-size:13.5pt; font-weight:700; color:#4b5563; white-space:nowrap; }}

  .card {{ flex:1; background:#fff; border:1px solid #e5e9f4; border-radius:12px;
           padding:.24in .32in; min-height:0; display:flex; flex-direction:column; line-height:1.45; }}
  .crow {{ display:flex; gap:.28in; min-height:0; }}
  .ccol {{ flex:1; min-width:0; }}
  .cond {{ background:#eef1fe; border-left:.065in solid #b9c2f5; border-radius:8px;
           padding:.09in .18in .10in; margin-bottom:.12in; color:#3f478f; }}
  .cond span {{ display:block; font-size:10.5pt; font-weight:700; margin-bottom:.03in; }}
  .miss-warn {{ background:#fffbeb; border:1px solid #fde68a; border-radius:8px; padding:.05in .14in;
                margin-bottom:.12in; font-size:12.5pt; font-weight:700; color:#b45309; }}
  .fallback {{ background:#fffbeb; border:1px solid #fde68a; border-radius:6px; padding:.04in .1in;
               margin-bottom:.08in; font-size:.62em; color:#b45309; }}
  .stem {{ font-size:1em; line-height:1.5; }}
  sup, sub {{ font-size:.6em; line-height:0; }}
  .opt {{ font-size:.85em; color:#4b5563; line-height:1.45; margin-top:.04in; }}
  .opt i {{ font-style:normal; color:#8b949e; }}
  .panel {{ width:{FIG_W}in; flex:none; border:1px solid #e5e9f4; border-radius:10px; background:#f2f4fb;
            display:flex; flex-direction:column; align-items:center; justify-content:center; padding:.14in; }}
  .panel img {{ max-width:100%; max-height:{BODY_H-0.66}in; }}
  .panel .cap {{ font-size:10pt; color:#8b949e; margin-top:.06in; }}
  .nofig {{ margin-top:.14in; padding:.06in .14in; border:1px dashed #f0d9a8; border-radius:8px;
            background:#fffbeb; color:#b45309; font-size:11.5pt; }}

  .ans {{ height:{ANS_H}in; border-radius:12px; padding:.09in .26in; display:flex; flex-direction:column; justify-content:center;
          background:#ebfbf3; border:1px solid #bfe9d6; overflow:hidden; margin-top:.14in; }}
  .ans .lab {{ font-size:11pt; color:#0f7b4f; font-weight:700; }}
  .ans .risk {{ color:#b45309; font-weight:400; }}
  .ans .val {{ font-weight:700; color:#0f7b4f; line-height:1.36; }}
  .ans.miss {{ background:#fffbeb; border-color:#fde68a; }}
  .ans.miss .lab {{ color:#b45309; }}
  .ans.miss .val {{ color:#b45309; font-size:18pt; }}
  .ans.plain {{ background:#fff; border-color:#e5e9f4; align-items:flex-start; }}
  .ans.plain .lab {{ color:#8b949e; }}
'''
    html = f'<!DOCTYPE html><html lang="zh-CN"><head><meta charset="utf-8"><title>{e(data.get("title"))}</title><style>{css}</style></head><body>{"".join(body)}</body></html>'
    out_html.write_text(html, encoding="utf-8")
    return out_html


def find_chrome():
    for p in CHROME_CANDIDATES:
        if os.path.exists(p):
            return p
    return None


def html_to_pdf(chrome, html_path: Path, pdf_path: Path):
    tmp = Path(tempfile.mkdtemp(prefix="mx_pdf_"))
    cmd = [chrome, "--headless=new", "--disable-gpu", "--no-sandbox", "--disable-extensions",
           "--no-pdf-header-footer", "--run-all-compositor-stages-before-draw",
           "--virtual-time-budget=20000",
           f"--user-data-dir={tmp}",
           f"--print-to-pdf={pdf_path}",
           html_path.as_uri()]
    try:
        r = subprocess.run(cmd, capture_output=True, timeout=300)
        if not pdf_path.exists() or pdf_path.stat().st_size == 0:
            # 老版 Chrome 回退
            cmd2 = [c for c in cmd if c != "--no-pdf-header-footer"] + ["--print-to-pdf-no-header"]
            cmd2[cmd2.index(f"--print-to-pdf={pdf_path}")] = f"--print-to-pdf={pdf_path}"
            subprocess.run(cmd2, capture_output=True, timeout=300)
        return pdf_path.exists() and pdf_path.stat().st_size > 0, (r.stderr or b"").decode("utf-8", "ignore")[-400:]
    except Exception as ex:
        return False, str(ex)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ────────────────────────────── 主流程 ──────────────────────────────
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("slides_json")
    ap.add_argument("--out", default=None)
    ap.add_argument("--no-animation", action="store_true", help="不出动画，答案直接可见")
    ap.add_argument("--page-image", action="store_true",
                    help="题目没有系统配图时，退回贴整页原卷页（默认关闭：投屏只给配图，原卷页太乱）")
    ap.add_argument("--no-page-image", action="store_true", help="（已废弃，等同于默认行为；保留兼容）")
    ap.add_argument("--no-pdf", action="store_true")
    ap.add_argument("--pdf-only", action="store_true")
    ap.add_argument("--cache", default=None)
    args = ap.parse_args()
    use_page_image = bool(args.page_image) and not args.no_page_image
    # [2026-09-21] --wb-image 已下线：整题裁片（wrong_questions.question_image_url，配图 B）
    # 在写入侧正式下线，slides.json 也不再下发 wbImage 字段。

    src = Path(args.slides_json).resolve()
    data = json.loads(src.read_text(encoding="utf-8"))
    out_dir = Path(args.out).resolve() if args.out else src.parent
    out_dir.mkdir(parents=True, exist_ok=True)
    with_answer = bool(data.get("withAnswer", True))

    # 文件名：直接沿用 slides.json 的基名，保证与 HTML/JSON 同名对齐
    name = src.stem[:-7] if src.stem.endswith(".slides") else src.stem

    cache = Path(args.cache).resolve() if args.cache else Path(tempfile.mkdtemp(prefix="mx_handout_img_"))
    cache.mkdir(parents=True, exist_ok=True)

    slides = data.get("slides", [])
    questions = [s for s in slides if s["kind"] == "question"]

    # ── 下载图片（默认只下配图；--page-image 才下原卷页）──
    need = []
    for q in questions:
        if q.get("figure"):
            need.append(q["figure"])
        elif use_page_image and FIG_HINT.search((q.get("parentStem") or "") + (q.get("stem") or "")):
            doc = next((st.get("docImage") for st in q.get("students", []) if st.get("docImage")), None)
            if doc:
                need.append(doc)
    need = sorted(set(need))
    print(f"[1] 需下载配图 {len(need)} 张（共 {len(questions)} 题）")
    imgmap_paths = {}
    if need:
        with concurrent.futures.ThreadPoolExecutor(max_workers=6) as ex:
            futs = {ex.submit(fetch_image, u, cache): u for u in need}
            for f in concurrent.futures.as_completed(futs):
                u = futs[f]
                imgmap_paths[u] = f.result()
        ok = sum(1 for v in imgmap_paths.values() if v)
        print(f"[2] 下载成功 {ok}/{len(need)}，缓存目录 {cache}")

    fig_for, page_for = {}, {}
    for q in questions:
        if q.get("figure"):
            fig_for[id(q)] = imgmap_paths.get(q["figure"])
        elif use_page_image and FIG_HINT.search((q.get("parentStem") or "") + (q.get("stem") or "")):
            doc = next((st.get("docImage") for st in q.get("students", []) if st.get("docImage")), None)
            page_for[id(q)] = imgmap_paths.get(doc) if doc else None

    # ── PPTX ──
    pptx_path = out_dir / (name + ".pptx")
    if not args.pdf_only:
        prs = Presentation()
        prs.slide_width = Inches(SLIDE_W)
        prs.slide_height = Inches(SLIDE_H)
        slide_cover(prs, data)
        slide_overview(prs, data)
        sec_no, total_sec = 0, data.get("stats", {}).get("sections", 0)
        for sl in slides:
            if sl["kind"] == "section":
                sec_no += 1
                slide_section(prs, sl, sec_no, total_sec)
            else:
                slide_question(prs, sl, fig_for.get(id(sl)), page_for.get(id(sl)),
                               animate=not args.no_animation, with_answer=with_answer)
        try:
            prs.save(str(pptx_path))
        except PermissionError:
            pptx_path = out_dir / (name + "_新.pptx")
            prs.save(str(pptx_path))
            print(f"    ! 原文件被占用（可能正开着），已改存：{pptx_path.name}")
        print(f"[3] PPTX 已生成：{pptx_path}  （{len(prs.slides)} 页）")

    # ── PDF ──
    pdf_path = out_dir / (name + ".pdf")
    if not args.no_pdf:
        html_path = out_dir / (name + ".print.html")
        imgmap = {id(q): (fig_for.get(id(q)), page_for.get(id(q))) for q in questions}
        build_print_html(data, slides, imgmap, html_path, with_answer)
        chrome = find_chrome()
        if chrome:
            ok, err = html_to_pdf(chrome, html_path, pdf_path)
            if ok:
                print(f"[4] PDF 已生成：{pdf_path}（{pdf_path.stat().st_size//1024} KB）")
            else:
                print(f"[4] ! PDF 生成失败：{err}")
        else:
            print("[4] ! 未找到 Chrome/Edge，跳过 PDF（可用浏览器打开 print.html 自行打印为 PDF）")

    summary = [
        f"课件：{data.get('title')}",
        f"时段：{data.get('period',{}).get('start')} ~ {data.get('period',{}).get('end')}",
        f"题目 {len(questions)} 题 / 分节 {data.get('stats',{}).get('sections')} 节 / 学生 {data.get('stats',{}).get('students')} 人",
        f"带配图（系统裁片）{sum(1 for q in questions if fig_for.get(id(q)))} 题；"
        f"原卷有图但无配图 {sum(1 for q in questions if not fig_for.get(id(q)) and not page_for.get(id(q)) and FIG_HINT.search((q.get('parentStem') or '') + (q.get('stem') or '')))} 题"
        + (f"；退贴原卷页 {sum(1 for q in questions if page_for.get(id(q)))} 题" if use_page_image else "（未贴原卷页）"),
        f"含答案 {'是' if with_answer else '否（无答案版）'}；答案{'单击浮现' if not args.no_animation and with_answer else '直接可见'}",
        f"PPTX：{pptx_path}",
        f"PDF ：{pdf_path if pdf_path.exists() else '（未生成）'}",
    ]
    (out_dir / (name + ".summary.txt")).write_text("\n".join(summary), encoding="utf-8")
    print("\n".join(summary))
    if not args.cache:
        print(f"（图片缓存 {cache}，可删）")


if __name__ == "__main__":
    main()
